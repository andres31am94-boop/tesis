"""
Genera los escenarios de demanda REAL a partir del aforo.

Para cada escenario de config/escenarios_aforo.yaml:
  1. Toma de la hoja "Aforo" las filas de esa fecha y periodo (debe haber exactamente 1 por brazo).
  2. Convierte cada conteo en vehiculos/hora: conteo * 60 / duracion_min.
  3. Escribe:
       sumo/demanda/<escenario>_flujos.csv   (trazabilidad: brazo, movimiento, tipo, vehiculos_hora)
       sumo/routes/<escenario>.rou.xml       (un flujo por brazo x movimiento x tipo)
       sumo/configs/<escenario>.sumocfg      (red + rutas + semilla)

A diferencia de la demanda de prueba, aqui cada brazo conserva SU PROPIA mezcla de tipos
y de giros, tal como se observo en campo.

Uso (desde la carpeta tesis):
  .\\.venv\\Scripts\\python.exe scripts\\generar_escenarios_aforo.py
"""
import csv
import sys
from pathlib import Path

import yaml
from openpyxl import load_workbook

from generar_demanda import MOVIMIENTOS, TIPOS_SUMO  # mismo directorio scripts/

CARPETA_TESIS = Path(__file__).resolve().parents[1]
PLANILLA = CARPETA_TESIS / "docs" / "aforo" / "planilla_aforo.xlsx"
CONFIG = CARPETA_TESIS / "config" / "escenarios_aforo.yaml"
CONFIG_CALIBRACION = CARPETA_TESIS / "config" / "calibracion.yaml"
DURACION_DEMANDA_S = 3600
BRAZOS = ["norte", "sur", "este", "oeste"]
TIPOS = ["carro", "moto", "bus", "camion"]
MOVS = ["recto", "derecha"]
# Columnas de conteo en la hoja Aforo (H..O): tipo x movimiento
COLUMNAS_CONTEO = [(t, m) for t in TIPOS for m in MOVS]

PLANTILLA_SUMOCFG = """<?xml version="1.0" encoding="UTF-8"?>
<!-- Escenario REAL generado por scripts/generar_escenarios_aforo.py: {descripcion} -->
<configuration>
    <input>
        <net-file value="../network/acacias.net.xml"/>
        <route-files value="../routes/{nombre}.rou.xml"/>
    </input>
    <time>
        <begin value="0"/>
        <end value="{duracion}"/>
    </time>
    <random_number>
        <seed value="42"/>
    </random_number>
    <report>
        <no-step-log value="true"/>
    </report>{procesamiento}
</configuration>
"""

PLANTILLA_SUBCARRILES = """
    <processing>
        <!-- Modelo de subcarriles (config/calibracion.yaml): permite a las motos ir lado a lado -->
        <lateral-resolution value="{resolucion}"/>
    </processing>"""


def leer_filas_aforo():
    hoja = load_workbook(PLANILLA, data_only=True)["Aforo"]
    filas = []
    for fila in hoja.iter_rows(min_row=2, values_only=True):
        if not fila[2]:
            continue
        brazo = fila[2].split()[0].lower()
        if brazo not in BRAZOS:
            sys.exit(f"ERROR: brazo desconocido en la planilla: {fila[2]}")
        conteos = fila[7:15]
        if any(c is None or c < 0 for c in conteos):
            sys.exit(f"ERROR: conteo vacio o negativo en la fila {fila}")
        filas.append({"fecha": fila[0].date(), "periodo": fila[3], "brazo": brazo,
                      "duracion_min": float(fila[6]), "conteos": dict(zip(COLUMNAS_CONTEO, conteos))})
    return filas


def flujos_de_escenario(filas, fecha, periodo):
    seleccion = [f for f in filas if f["fecha"] == fecha and f["periodo"] == periodo]
    brazos = sorted(f["brazo"] for f in seleccion)
    if brazos != sorted(BRAZOS):
        sys.exit(f"ERROR: la sesion {fecha} {periodo} debe tener exactamente un conteo por brazo; tiene {brazos}")
    flujos = []
    for f in seleccion:
        for (tipo, mov), conteo in f["conteos"].items():
            flujos.append((f["brazo"], mov, tipo, conteo * 60.0 / f["duracion_min"]))
    return flujos


def escribir_rutas(nombre, flujos, descripcion, atributos_moto, atributos_flujo_moto, usar_accesos):
    lineas = ['<?xml version="1.0" encoding="UTF-8"?>',
              "<!-- Generado por scripts/generar_escenarios_aforo.py: NO editar a mano. -->",
              f"<!-- Escenario {nombre}: {descripcion}. Fuente: docs/aforo/planilla_aforo.xlsx -->",
              "<routes>"]
    for tipo in TIPOS:
        atributos = TIPOS_SUMO[tipo]
        extra = "".join(f' {clave}="{valor}"' for clave, valor in atributos_moto.items()) if tipo == "moto" else ""
        lineas.append(f'    <vType id="{tipo}" vClass="{atributos["vClass"]}" guiShape="{atributos["guiShape"]}"{extra}/>')
    for brazo, mov, tipo, vehiculos_hora in flujos:
        if vehiculos_hora == 0:
            continue
        probabilidad = vehiculos_hora / 3600.0
        if probabilidad > 1:
            sys.exit(f"ERROR: {brazo}/{mov}/{tipo} supera 3600 veh/h")
        llegada, salida = MOVIMIENTOS[(brazo, mov)]
        if usar_accesos:
            llegada = f"acceso_{brazo}"  # tramo creado por scripts/generar_accesos.py
        extra = "".join(f' {clave}="{valor}"' for clave, valor in atributos_flujo_moto.items()) if tipo == "moto" else ""
        lineas.append(f'    <flow id="{brazo}_{mov}_{tipo}" type="{tipo}" begin="0" end="{DURACION_DEMANDA_S}" '
                      f'from="{llegada}" to="{salida}" probability="{probabilidad:.6f}" '
                      f'departLane="best" departSpeed="max"{extra}/>')
    lineas.append("</routes>")
    (CARPETA_TESIS / "sumo" / "routes" / f"{nombre}.rou.xml").write_text("\n".join(lineas) + "\n", encoding="utf-8")


def main():
    with open(CONFIG, encoding="utf-8") as archivo:
        escenarios = yaml.safe_load(archivo)["escenarios"]
    with open(CONFIG_CALIBRACION, encoding="utf-8") as archivo:
        calibracion = yaml.safe_load(archivo)
    subcarriles = calibracion["subcarriles"]
    procesamiento = (PLANTILLA_SUBCARRILES.format(resolucion=subcarriles["resolucion_lateral_m"])
                     if subcarriles["activo"] else "")
    atributos_moto = calibracion.get("moto", {}) if subcarriles["activo"] else {}
    atributos_flujo_moto = calibracion.get("moto_flujo", {}) if subcarriles["activo"] else {}
    usar_accesos = calibracion.get("accesos", {}).get("activo", False)
    print(f"Subcarriles: {'SI' if subcarriles['activo'] else 'NO'} | atributos moto: {atributos_moto} | "
          f"flujos de motos: {atributos_flujo_moto} | accesos: {'SI' if usar_accesos else 'NO'}")
    filas = leer_filas_aforo()
    for nombre, datos in escenarios.items():
        flujos = flujos_de_escenario(filas, datos["fecha"], datos["periodo"])
        with open(CARPETA_TESIS / "sumo" / "demanda" / f"{nombre}_flujos.csv", "w", newline="", encoding="utf-8") as archivo:
            escritor = csv.writer(archivo)
            escritor.writerow(["brazo_llegada", "movimiento", "tipo", "vehiculos_hora"])
            escritor.writerows(flujos)
        escribir_rutas(nombre, flujos, datos["descripcion"], atributos_moto, atributos_flujo_moto, usar_accesos)
        (CARPETA_TESIS / "sumo" / "configs" / f"{nombre}.sumocfg").write_text(
            PLANTILLA_SUMOCFG.format(nombre=nombre, descripcion=datos["descripcion"], duracion=DURACION_DEMANDA_S,
                                     procesamiento=procesamiento),
            encoding="utf-8")
        por_brazo = {b: sum(v for br, _, _, v in flujos if br == b) for b in BRAZOS}
        print(f"{nombre:<20} total {sum(por_brazo.values()):6.0f} veh/h | " +
              " | ".join(f"{b} {v:5.0f}" for b, v in por_brazo.items()))


if __name__ == "__main__":
    main()
